"""tabelle_in_db.py — le tabelle CSV o Excel di una cartella in un database SQLite.

Una tabella per file, con il nome del file; le colonne sono quelle della
prima riga. Rilanciato, rifà le tabelle da capo. Si lancia così com'è, oppure
se ne copiano le righe nel proprio file.

Uso (dalla cartella scripts/):
    python scaffolding/tabelle_in_db.py dati/tabelle dati/tabelle.db
    python scaffolding/tabelle_in_db.py temi/valtesa/tabelle temi/valtesa/valtesa.db

I CSV vanno con la virgola tra le colonne e il punto nei decimali. Se i dati
sono in Excel, il file .xlsx si legge così com'è.
"""
import csv
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

cartella, database = Path(sys.argv[1]), sys.argv[2]
con = sqlite3.connect(database)


def tipo(valori):
    """INTEGER se sono tutti interi, REAL se sono tutti numeri, altrimenti TEXT (le celle vuote non contano)."""
    pieni = [str(v) for v in valori if v is not None]
    for nome, converti in (("INTEGER", int), ("REAL", float)):
        try:
            [converti(v) for v in pieni]
            return nome
        except ValueError:
            pass
    return "TEXT"


def salva(nome, intestazione, righe):
    """Una tabella nel database: CREATE TABLE, poi un INSERT per riga."""
    tipi = [tipo([r[i] for r in righe]) for i in range(len(intestazione))]
    colonne = ", ".join(f'"{c}" {t}' for c, t in zip(intestazione, tipi))
    segnaposto = ", ".join("?" for c in intestazione)
    con.execute(f'DROP TABLE IF EXISTS "{nome}"')
    con.execute(f'CREATE TABLE "{nome}" ({colonne})')
    con.executemany(f'INSERT INTO "{nome}" VALUES ({segnaposto})', righe)
    print(f"{nome}: {len(righe)} righe")


for file in sorted(cartella.glob("*.csv")):
    with open(file, newline="", encoding="utf-8-sig") as f:
        intestazione, *righe = csv.reader(f)
    salva(file.stem, intestazione, [[v or None for v in r] for r in righe])   # cella vuota: NULL

for file in sorted(cartella.glob("*.xlsx")):
    intestazione, *righe = load_workbook(file).active.iter_rows(values_only=True)
    salva(file.stem, intestazione,
          [[v.date().isoformat() if isinstance(v, datetime) else v for v in r] for r in righe])   # le date: 2026-07-15

con.commit()
con.close()
